"""Deterministic loaders for the ERP goods-received CSV and the store sheet."""

from __future__ import annotations

import csv
import math
import re
import warnings
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.workbook.workbook import Workbook
from pydantic import BaseModel

from src.config import KARACHI, RulesConfig, load_rules

_GRN_PREFIX = re.compile(r"(?i)^grn-")
_US_DATE = "%m/%d/%Y"
_ISO_DATE = "%Y-%m-%d"
_DMY_DATE = "%d/%m/%Y"

_UOM_SINGULAR = {
    "PC": "PCS",
    "PCS": "PCS",
    "BAG": "BAG",
    "BAGS": "BAG",
    "CONE": "CONE",
    "CONES": "CONE",
    "DRUM": "DRUM",
    "DRUMS": "DRUM",
    "ROLL": "ROLL",
    "ROLLS": "ROLL",
    "CTN": "CTN",
    "CTNS": "CTN",
    "BDL": "BDL",
    "BDLS": "BDL",
    "YD": "YDS",
    "YDS": "YDS",
    "M": "M",
    "METER": "M",
    "METERS": "M",
    "METRE": "M",
    "METRES": "M",
}


class DataFormatWarning(UserWarning):
    """Raised when a store-sheet cell is text where a datetime was expected."""


class ErpGrnRow(BaseModel):
    source_pointer: str
    grn_no: str
    grn_date: str
    supplier_code: str
    supplier_name: str
    item_desc: str
    lot_no: str
    qty: float
    uom: str
    entered_by: str
    entered_at: datetime
    is_late_entry: bool
    late_entry_reason: str | None = None


class StoreSheetRow(BaseModel):
    source_pointer: str
    sheet_row: int
    grn_no: str | None
    entity_key: str
    receipt_date: str
    supplier: str
    material: str
    lot: str
    qty: float
    uom: str
    remarks: str | None = None
    is_duplicate: bool = False
    is_return: bool = False
    date_warning: str | None = None


def load_erp_grn(path: Path | str, rules: RulesConfig | None = None) -> list[ErpGrnRow]:
    rules = rules or load_rules()
    trusted = {user.strip().lower() for user in rules.audit_rules.trusted_erp_users}
    max_delay = timedelta(hours=rules.audit_rules.max_entry_delay_hours)
    rows: list[ErpGrnRow] = []
    with Path(path).open(newline="", encoding="utf-8") as handle:
        reader = csv.DictReader(handle)
        for index, raw in enumerate(reader):
            grn_no = canonical_grn(raw["grn_no"])
            if grn_no is None:
                raise ValueError(f"ERP data row {index + 1} has no GRN number")
            grn_date = _iso_date(raw["grn_date"].strip())
            entered_at = _parse_entered_at(raw["entered_at"].strip())
            entered_by = raw["entered_by"].strip()
            is_late, reason = _late_entry(
                entered_by=entered_by,
                entered_at=entered_at,
                grn_date=grn_date,
                trusted=trusted,
                max_delay=max_delay,
            )
            rows.append(
                ErpGrnRow(
                    source_pointer=f"[ERP R{index + 1:02d}]",
                    grn_no=grn_no,
                    grn_date=grn_date,
                    supplier_code=raw["supplier_code"].strip(),
                    supplier_name=raw["supplier_name"].strip(),
                    item_desc=raw["item_desc"].strip(),
                    lot_no=raw["lot_no"].strip(),
                    qty=float(raw["qty"]),
                    uom=canonical_uom(raw["uom"]),
                    entered_by=entered_by,
                    entered_at=entered_at,
                    is_late_entry=is_late,
                    late_entry_reason=reason,
                )
            )
    return rows


def load_store_sheet(path: Path | str, sheet_name: str = "Sep-2026") -> list[StoreSheetRow]:
    workbook = load_workbook(Path(path), data_only=False)
    try:
        return _read_store_sheet(workbook, sheet_name)
    finally:
        workbook.close()


def _read_store_sheet(workbook: Workbook, sheet_name: str) -> list[StoreSheetRow]:
    if sheet_name not in workbook.sheetnames:
        raise ValueError(f"Workbook has no sheet {sheet_name!r}")
    worksheet = workbook[sheet_name]
    header = [cell.value for cell in next(worksheet.iter_rows(min_row=1, max_row=1))]
    expected = ["Date", "GRN", "Supplier", "Material", "Lot", "Qty Received", "UOM", "Remarks"]
    if header[:8] != expected:
        raise ValueError(f"Unexpected store-sheet header: {header[:8]!r}")

    parsed: list[StoreSheetRow] = []
    seen_grn: set[str] = set()
    for cells in worksheet.iter_rows(min_row=2, max_row=worksheet.max_row, max_col=8):
        if _row_is_blank(cells):
            continue
        sheet_row = cells[0].row
        grn_no = canonical_grn(cells[1].value)
        receipt_date, date_warning = _parse_sheet_date(cells[0], sheet_row, grn_no)
        supplier = _text(cells[2].value)
        material = _text(cells[3].value)
        lot = _text(cells[4].value)
        qty = float(cells[5].value)
        uom = _text(cells[6].value)
        remarks = _text(cells[7].value) or None
        is_return = grn_no is None
        entity_key = f"RETURN-{lot}" if is_return else grn_no
        if entity_key is None:
            raise ValueError(f"Excel row {sheet_row} has neither a GRN nor a lot")
        is_duplicate = grn_no is not None and grn_no in seen_grn
        if grn_no is not None and not is_duplicate:
            seen_grn.add(grn_no)
        parsed.append(
            StoreSheetRow(
                source_pointer=f"[Excel R{sheet_row:02d}]",
                sheet_row=sheet_row,
                grn_no=grn_no,
                entity_key=entity_key,
                receipt_date=receipt_date,
                supplier=supplier,
                material=material,
                lot=lot,
                qty=qty,
                uom=uom,
                remarks=remarks,
                is_duplicate=is_duplicate,
                is_return=is_return,
                date_warning=date_warning,
            )
        )
    return parsed


def canonical_grn(value: object) -> str | None:
    if value is None:
        return None
    if isinstance(value, float) and math.isnan(value):
        return None
    if isinstance(value, float) and value.is_integer():
        text = str(int(value))
    elif isinstance(value, int):
        text = str(value)
    else:
        text = str(value).strip()
    if not text or text.lower() in {"nan", "none"}:
        return None
    text = _GRN_PREFIX.sub("", text).strip()
    if text.endswith(".0") and text[:-2].isdigit():
        text = text[:-2]
    return text


def canonical_uom(value: object) -> str:
    token = str(value).strip().upper()
    return _UOM_SINGULAR.get(token, token)


def _iso_date(value: str) -> str:
    return datetime.strptime(value, _ISO_DATE).date().isoformat()


def _parse_entered_at(value: str) -> datetime:
    parsed = datetime.strptime(value, "%Y-%m-%d %H:%M")
    return parsed.replace(tzinfo=KARACHI)


def _late_entry(
    *,
    entered_by: str,
    entered_at: datetime,
    grn_date: str,
    trusted: set[str],
    max_delay: timedelta,
) -> tuple[bool, str | None]:
    reasons: list[str] = []
    if entered_by.lower() not in trusted:
        reasons.append(f"entered_by {entered_by!r} is not a trusted ERP user")
    receipt_start = datetime.strptime(grn_date, _ISO_DATE).replace(tzinfo=KARACHI)
    delay = entered_at - receipt_start
    if delay > max_delay:
        hours = delay.total_seconds() / 3600
        reasons.append(f"entry delay {hours:.2f}h exceeds {max_delay.total_seconds() / 3600:g}h")
    if not reasons:
        return False, None
    return True, "; ".join(reasons)


def _parse_sheet_date(cell: Cell, sheet_row: int, grn_no: str | None) -> tuple[str, str | None]:
    value = cell.value
    if isinstance(value, datetime):
        return value.date().isoformat(), None
    if isinstance(value, date):
        return value.isoformat(), None
    if isinstance(value, str):
        iso = _parse_text_date(value.strip())
        label = grn_no or "blank GRN"
        message = (
            f"DataFormatWarning: Excel row {sheet_row} ({label}) date {value!r} "
            f"is text, not a datetime; normalized to {iso}"
        )
        warnings.warn(message, DataFormatWarning, stacklevel=2)
        return iso, message
    raise ValueError(f"Excel row {sheet_row} has an unsupported date cell {value!r}")


def _parse_text_date(value: str) -> str:
    for fmt in (_US_DATE, _ISO_DATE, _DMY_DATE):
        try:
            return datetime.strptime(value, fmt).date().isoformat()
        except ValueError:
            continue
    raise ValueError(f"Cannot parse store-sheet date {value!r}")


def _text(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _row_is_blank(cells: tuple[Cell, ...]) -> bool:
    return all(cell.value is None or str(cell.value).strip() == "" for cell in cells)
